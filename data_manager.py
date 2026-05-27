# data_manager.py — SQLite-backed data layer for Florin
import uuid
import json
from pathlib import Path
from datetime import datetime
import msoffcrypto
import openpyxl
import io
from db import get_connection
from theme import CATEGORY_COLORS, COLOR_TEXT_FAINT

DATA_DIR = Path("data")
CONFIG_FILE = Path("config.json")


def generate_id():
    return str(uuid.uuid4())


def init_env():
    """Initialize database and run migration if needed."""
    import migrate
    migrate.run_migration()


# --- Settings ---
def get_setting(key, default=None):
    conn = get_connection()
    row = conn.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
    return row["value"] if row else default


def set_setting(key, value):
    conn = get_connection()
    conn.execute("INSERT OR REPLACE INTO settings (key, value) VALUES (?, ?)", (key, str(value)))
    conn.commit()


def get_config():
    """Return config dict compatible with old API."""
    conn = get_connection()
    config = {
        "last_month": get_setting("last_month", datetime.now().strftime("%Y-%m")),
        "theme": get_setting("theme", "light"),
        "language": get_setting("language", "en"),
    }
    config["cashflow_targets"] = get_cashflow_targets()
    config["default_income_items"] = get_default_income_items()
    return config


def save_config(config):
    """Persist config dict (compatibility shim)."""
    for key in ("last_month", "theme", "language"):
        if key in config:
            set_setting(key, config[key])
    if "cashflow_targets" in config:
        conn = get_connection()
        conn.execute("DELETE FROM cashflow_targets")
        for i, t in enumerate(config["cashflow_targets"]):
            conn.execute(
                "INSERT INTO cashflow_targets (id, name, target, sort_order) VALUES (?, ?, ?, ?)",
                (t["id"], t["name"], t["target"], i)
            )
        conn.commit()
    if "default_income_items" in config:
        conn = get_connection()
        conn.execute("DELETE FROM default_income_items")
        for i, item in enumerate(config["default_income_items"]):
            conn.execute(
                "INSERT INTO default_income_items (id, name, type, sort_order) VALUES (?, ?, ?, ?)",
                (item["id"], item["name"], item["type"], i)
            )
        conn.commit()



# --- Default Income Items ---
def get_default_income_items(config=None):
    conn = get_connection()
    rows = conn.execute("SELECT id, name, type FROM default_income_items ORDER BY sort_order").fetchall()
    return [{"id": r["id"], "name": r["name"], "type": r["type"]} for r in rows]


def add_default_income_item(name, item_type="addition"):
    conn = get_connection()
    item_id = generate_id()
    max_order = conn.execute("SELECT COALESCE(MAX(sort_order), -1) FROM default_income_items").fetchone()[0]
    conn.execute(
        "INSERT INTO default_income_items (id, name, type, sort_order) VALUES (?, ?, ?, ?)",
        (item_id, name, item_type, max_order + 1)
    )
    conn.commit()
    return {"id": item_id, "name": name, "type": item_type}


def delete_default_income_item(item_id):
    conn = get_connection()
    conn.execute("DELETE FROM default_income_items WHERE id = ?", (item_id,))
    conn.commit()


# --- Cashflow Targets ---
def get_cashflow_targets(config=None):
    conn = get_connection()
    rows = conn.execute("SELECT id, name, target FROM cashflow_targets ORDER BY sort_order").fetchall()
    return [{"id": r["id"], "name": r["name"], "target": r["target"]} for r in rows]


def add_cashflow_target(name, target_amount):
    conn = get_connection()
    tid = generate_id()
    max_order = conn.execute("SELECT COALESCE(MAX(sort_order), -1) FROM cashflow_targets").fetchone()[0]
    conn.execute(
        "INSERT INTO cashflow_targets (id, name, target, sort_order) VALUES (?, ?, ?, ?)",
        (tid, name, target_amount, max_order + 1)
    )
    conn.commit()
    return {"id": tid, "name": name, "target": target_amount}


def edit_cashflow_target(target_id, name=None, target_amount=None):
    conn = get_connection()
    if name is not None:
        conn.execute("UPDATE cashflow_targets SET name = ? WHERE id = ?", (name, target_id))
    if target_amount is not None:
        conn.execute("UPDATE cashflow_targets SET target = ? WHERE id = ?", (target_amount, target_id))
    conn.commit()
    row = conn.execute("SELECT id, name, target FROM cashflow_targets WHERE id = ?", (target_id,)).fetchone()
    return {"id": row["id"], "name": row["name"], "target": row["target"]} if row else None


def delete_cashflow_target(target_id):
    conn = get_connection()
    conn.execute("DELETE FROM cashflow_targets WHERE id = ?", (target_id,))
    conn.commit()


# --- Month Data ---
def _get_previous_month_str(month_str):
    year, month = int(month_str[:4]), int(month_str[5:7])
    month -= 1
    if month == 0:
        month = 12
        year -= 1
    return f"{year:04d}-{month:02d}"


def get_month_filepath(month_str):
    return DATA_DIR / f"{month_str}.json"


def get_past_months(current_month):
    """Return list of month strings that have data, excluding current."""
    conn = get_connection()
    rows = conn.execute(
        "SELECT DISTINCT month FROM income_items WHERE month != ? UNION SELECT DISTINCT month FROM expenses WHERE month != ? ORDER BY month DESC",
        (current_month, current_month)
    ).fetchall()
    return [r[0] for r in rows]


def save_snapshot(month_str, data):
    """Persist an immutable JSON snapshot of a month and return its metadata."""
    conn = get_connection()
    snapshot_id = generate_id()
    payload = json.dumps(data, ensure_ascii=False, sort_keys=True)
    conn.execute(
        "INSERT INTO snapshots (id, month, payload) VALUES (?, ?, ?)",
        (snapshot_id, month_str, payload)
    )
    conn.commit()
    row = conn.execute(
        "SELECT id, month, created_at FROM snapshots WHERE id = ?",
        (snapshot_id,)
    ).fetchone()
    return {"id": row["id"], "month": row["month"], "created_at": row["created_at"]}


def get_snapshots(month_str=None):
    conn = get_connection()
    if month_str:
        rows = conn.execute(
            "SELECT id, month, created_at FROM snapshots WHERE month = ? ORDER BY created_at DESC",
            (month_str,)
        ).fetchall()
    else:
        rows = conn.execute(
            "SELECT id, month, created_at FROM snapshots ORDER BY created_at DESC"
        ).fetchall()
    return [{"id": r["id"], "month": r["month"], "created_at": r["created_at"]} for r in rows]



DEFAULT_EXPENSE_CATEGORIES = [
    {"id": "groceries", "name": "Groceries"},
    {"id": "healthcare", "name": "Healthcare"},
    {"id": "beauty", "name": "Beauty"},
    {"id": "transport", "name": "Transport"},
    {"id": "subscriptions", "name": "Subscriptions"},
    {"id": "household", "name": "Household"},
    {"id": "eating_out", "name": "Eating Out"},
    {"id": "other", "name": "Other"},
]


def load_month(month_str):
    """Load month data from SQLite into dict (same shape as old JSON)."""
    init_env()
    conn = get_connection()
    # Check if month exists
    row = conn.execute("SELECT 1 FROM categories WHERE month = ? LIMIT 1", (month_str,)).fetchone()
    if row is None:
        data = generate_default_month(month_str)
        save_month(month_str, data)
        return data
    return _load_month_from_db(conn, month_str)


def _load_month_from_db(conn, month_str):
    """Reconstruct month dict from SQLite tables."""
    data = {"month": month_str}

    # Categories
    rows = conn.execute("SELECT id, name, percent, color FROM categories WHERE month = ?", (month_str,)).fetchall()
    data["categories"] = [{"id": r["id"], "name": r["name"], "percent": r["percent"], "color": r["color"]} for r in rows]

    # Expense categories
    rows = conn.execute("SELECT id, name FROM expense_categories WHERE month = ?", (month_str,)).fetchall()
    data["expense_categories"] = [{"id": r["id"], "name": r["name"]} for r in rows]

    # Income items
    rows = conn.execute("SELECT id, name, amount, type, category_id FROM income_items WHERE month = ?", (month_str,)).fetchall()
    data["income_items"] = [{"id": r["id"], "name": r["name"], "amount": r["amount"], "type": r["type"], "category_id": r["category_id"]} for r in rows]

    # Expenses
    rows = conn.execute("SELECT id, date, amount, category_id, expense_category_id, savings_category_id, description, tags FROM expenses WHERE month = ?", (month_str,)).fetchall()
    data["expenses"] = [{
        "id": r["id"], "date": r["date"], "amount": r["amount"],
        "category_id": r["category_id"], "expense_category_id": r["expense_category_id"],
        "savings_category_id": r["savings_category_id"], "description": r["description"],
        "tags": [t for t in r["tags"].split(",") if t]
    } for r in rows]

    # Cashflow
    cf_rows = conn.execute("SELECT key, value, updated_at FROM cashflow WHERE month = ?", (month_str,)).fetchall()
    cashflow = {"current_accounts": {}, "shared_actual": 0.0, "shared_assumed": 0.0, "saved_assumed": 0.0}
    for r in cf_rows:
        if r["key"].startswith("acct:"):
            acct_id = r["key"][5:]
            cashflow["current_accounts"][acct_id] = {"balance": r["value"], "updated_at": r["updated_at"]}
        else:
            cashflow[r["key"]] = r["value"]
    data["cashflow"] = cashflow

    return data


def save_month(month_str, data):
    """Persist month dict to SQLite."""
    conn = get_connection()
    # Categories
    conn.execute("DELETE FROM categories WHERE month = ?", (month_str,))
    for cat in data.get("categories", []):
        conn.execute(
            "INSERT INTO categories (id, month, name, percent, color) VALUES (?, ?, ?, ?, ?)",
            (cat["id"], month_str, cat["name"], cat.get("percent", 0), cat.get("color", CATEGORY_COLORS.get(cat["id"], COLOR_TEXT_FAINT)))
        )

    # Expense categories
    conn.execute("DELETE FROM expense_categories WHERE month = ?", (month_str,))
    for cat in data.get("expense_categories", []):
        conn.execute(
            "INSERT INTO expense_categories (id, month, name) VALUES (?, ?, ?)",
            (cat["id"], month_str, cat["name"])
        )

    # Income items
    conn.execute("DELETE FROM income_items WHERE month = ?", (month_str,))
    for item in data.get("income_items", []):
        conn.execute(
            "INSERT INTO income_items (id, month, name, amount, type, category_id) VALUES (?, ?, ?, ?, ?, ?)",
            (item["id"], month_str, item["name"], item.get("amount", 0.0), item["type"], item.get("category_id"))
        )

    # Expenses
    conn.execute("DELETE FROM expenses WHERE month = ?", (month_str,))
    for exp in data.get("expenses", []):
        tags = ",".join(exp.get("tags", []))
        conn.execute(
            "INSERT INTO expenses (id, month, date, amount, category_id, expense_category_id, savings_category_id, description, tags) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (exp["id"], month_str, exp["date"], exp["amount"], exp.get("category_id"), exp.get("expense_category_id"), exp.get("savings_category_id"), exp.get("description", ""), tags)
        )

    # Cashflow
    conn.execute("DELETE FROM cashflow WHERE month = ?", (month_str,))
    cf = data.get("cashflow", {})
    for key, val in cf.items():
        if key == "current_accounts":
            for acct_id, acct in val.items():
                if isinstance(acct, dict):
                    conn.execute(
                        "INSERT INTO cashflow (month, key, value, updated_at) VALUES (?, ?, ?, ?)",
                        (month_str, f"acct:{acct_id}", acct.get("balance", 0.0), acct.get("updated_at"))
                    )
                else:
                    conn.execute("INSERT INTO cashflow (month, key, value) VALUES (?, ?, ?)", (month_str, f"acct:{acct_id}", float(acct)))
        elif isinstance(val, (int, float)):
            conn.execute("INSERT INTO cashflow (month, key, value) VALUES (?, ?, ?)", (month_str, key, float(val)))
    conn.commit()
    # Update last_month setting
    set_setting("last_month", month_str)


def generate_default_month(month_str):
    """Generate default month data, carrying over from previous month if available."""
    conn = get_connection()
    prev = _get_previous_month_str(month_str)
    prev_exists = conn.execute("SELECT 1 FROM categories WHERE month = ? LIMIT 1", (prev,)).fetchone()

    if prev_exists:
        prev_data = _load_month_from_db(conn, prev)
        categories = prev_data.get("categories", _default_categories())
        expense_categories = prev_data.get("expense_categories", list(DEFAULT_EXPENSE_CATEGORIES))
    else:
        categories = _default_categories()
        expense_categories = list(DEFAULT_EXPENSE_CATEGORIES)

    default_items = get_default_income_items()
    income_items = [
        {"id": generate_id(), "name": di["name"], "amount": 0.0, "type": di["type"], "category_id": None}
        for di in default_items
    ]

    return {
        "month": month_str,
        "income_items": income_items,
        "categories": categories,
        "expense_categories": expense_categories,
        "expenses": [],
        "cashflow": {"current_accounts": {}, "shared_actual": 0.0, "shared_assumed": 0.0, "saved_assumed": 0.0}
    }


def _default_categories():
    return [
        {"id": "daily_life", "name": "Daily Life", "percent": 40, "color": CATEGORY_COLORS["daily_life"]},
        {"id": "shared", "name": "Shared", "percent": 20, "color": CATEGORY_COLORS["shared"]},
        {"id": "saved", "name": "Saved", "percent": 20, "color": CATEGORY_COLORS["saved"]},
        {"id": "pleasure", "name": "Pleasure", "percent": 20, "color": CATEGORY_COLORS["pleasure"]},
    ]



# --- Savings Planner ---
def get_savings_planner():
    """Load savings planner as dict (same shape as old JSON)."""
    init_env()
    conn = get_connection()
    sp = {"year": datetime.now().year, "categories": [], "groups": [], "grid": {}, "assumed": {}}

    # Groups
    rows = conn.execute("SELECT id, name FROM savings_groups").fetchall()
    sp["groups"] = [{"id": r["id"], "name": r["name"]} for r in rows]

    # Categories
    rows = conn.execute("SELECT id, name, target, deadline_month, group_id, next_year_target FROM savings_categories").fetchall()
    sp["categories"] = [{
        "id": r["id"], "name": r["name"], "target": r["target"],
        "deadline_month": r["deadline_month"], "group_id": r["group_id"],
        "next_year_target": r["next_year_target"]
    } for r in rows]

    # Grid
    rows = conn.execute("SELECT category_id, month_key, amount FROM savings_grid").fetchall()
    for r in rows:
        sp["grid"].setdefault(r["category_id"], {})[r["month_key"]] = r["amount"]

    # Assumed
    rows = conn.execute("SELECT month_key, amount FROM savings_assumed").fetchall()
    sp["assumed"] = {r["month_key"]: r["amount"] for r in rows}

    return sp


def save_savings_planner(sp):
    """Persist savings planner dict to SQLite."""
    conn = get_connection()

    # Groups
    conn.execute("DELETE FROM savings_groups")
    for g in sp.get("groups", []):
        conn.execute("INSERT INTO savings_groups (id, name) VALUES (?, ?)", (g["id"], g["name"]))

    # Categories
    conn.execute("DELETE FROM savings_categories")
    for cat in sp.get("categories", []):
        conn.execute(
            "INSERT INTO savings_categories (id, name, target, deadline_month, group_id, next_year_target) VALUES (?, ?, ?, ?, ?, ?)",
            (cat["id"], cat["name"], cat.get("target", 0), cat.get("deadline_month", 12), cat.get("group_id"), cat.get("next_year_target"))
        )

    # Grid
    conn.execute("DELETE FROM savings_grid")
    for cat_id, months in sp.get("grid", {}).items():
        for mk, amt in months.items():
            if amt:
                conn.execute("INSERT INTO savings_grid (category_id, month_key, amount) VALUES (?, ?, ?)", (cat_id, mk, amt))

    # Assumed
    conn.execute("DELETE FROM savings_assumed")
    for mk, amt in sp.get("assumed", {}).items():
        if amt:
            conn.execute("INSERT INTO savings_assumed (month_key, amount) VALUES (?, ?)", (mk, amt))

    # Actual grid
    conn.execute("DELETE FROM savings_actual_grid")
    for cat_id, months in sp.get("actual_grid", {}).items():
        for mk, amt in months.items():
            if amt:
                conn.execute("INSERT INTO savings_actual_grid (category_id, month_key, amount) VALUES (?, ?, ?)", (cat_id, mk, amt))

    # Actual available
    conn.execute("DELETE FROM savings_actual_available")
    for mk, amt in sp.get("actual_available", {}).items():
        if amt:
            conn.execute("INSERT INTO savings_actual_available (month_key, amount) VALUES (?, ?)", (mk, amt))

    # Spent
    conn.execute("DELETE FROM savings_spent")
    for cat_id, amt in sp.get("spent", {}).items():
        if amt:
            conn.execute("INSERT INTO savings_spent (category_id, amount) VALUES (?, ?)", (cat_id, amt))

    conn.commit()


def get_savings_actual(sp):
    """Load actual tracking data into sp dict."""
    conn = get_connection()
    sp.setdefault("actual_grid", {})
    sp.setdefault("actual_available", {})
    sp.setdefault("spent", {})

    rows = conn.execute("SELECT category_id, month_key, amount FROM savings_actual_grid").fetchall()
    for r in rows:
        sp["actual_grid"].setdefault(r["category_id"], {})[r["month_key"]] = r["amount"]

    rows = conn.execute("SELECT month_key, amount FROM savings_actual_available").fetchall()
    for r in rows:
        sp["actual_available"][r["month_key"]] = r["amount"]

    rows = conn.execute("SELECT category_id, amount FROM savings_spent").fetchall()
    for r in rows:
        sp["spent"][r["category_id"]] = r["amount"]

    return sp


def add_savings_category(sp, name, target, deadline_month=12, group_id=None, next_year_target=None):
    cat = {
        "id": generate_id(), "name": name, "target": target,
        "deadline_month": deadline_month, "group_id": group_id, "next_year_target": next_year_target,
    }
    sp["categories"].append(cat)
    months_count = deadline_month
    monthly = round(target / months_count, 2) if months_count > 0 else 0
    grid_row = {f"{m:02d}": monthly if m <= deadline_month else 0.0 for m in range(1, 13)}
    if next_year_target and deadline_month < 12:
        remaining_months = 12 - deadline_month
        monthly_nyt = round(next_year_target / remaining_months, 2)
        for m in range(deadline_month + 1, 13):
            grid_row[f"{m:02d}"] = monthly_nyt
    sp["grid"][cat["id"]] = grid_row
    return cat


def edit_savings_category(sp, cat_id, **fields):
    for cat in sp["categories"]:
        if cat["id"] == cat_id:
            cat.update(fields)
            return cat
    return None


def delete_savings_category(sp, cat_id):
    sp["categories"] = [c for c in sp["categories"] if c["id"] != cat_id]
    sp["grid"].pop(cat_id, None)


def add_savings_group(sp, name):
    g = {"id": generate_id(), "name": name}
    sp["groups"].append(g)
    return g


def edit_savings_group(sp, group_id, name):
    for g in sp["groups"]:
        if g["id"] == group_id:
            g["name"] = name
            return g
    return None


def delete_savings_group(sp, group_id):
    sp["groups"] = [g for g in sp["groups"] if g["id"] != group_id]
    for cat in sp["categories"]:
        if cat.get("group_id") == group_id:
            cat["group_id"] = None



# --- Savings Calculations ---
def get_month_total_allocated(sp, month_key):
    total = 0.0
    for cat_id, months in sp.get("grid", {}).items():
        total += months.get(month_key, 0.0)
    return total


def get_month_remaining(sp, month_key):
    assumed = sp.get("assumed", {}).get(month_key, 0.0)
    return assumed - get_month_total_allocated(sp, month_key)


def get_actual_month_total(sp, month_key):
    total = 0.0
    for cat_id, months in sp.get("actual_grid", {}).items():
        total += months.get(month_key, 0.0)
    return total


def get_actual_month_remaining(sp, month_key):
    available = sp.get("actual_available", {}).get(month_key, 0.0)
    return available - get_actual_month_total(sp, month_key)


def get_cat_total_saved(sp, cat_id):
    return sum(sp.get("actual_grid", {}).get(cat_id, {}).values())


def get_cat_missing(sp, cat_id):
    cat = next((c for c in sp.get("categories", []) if c["id"] == cat_id), None)
    if not cat:
        return 0.0
    return max(0.0, cat["target"] - get_cat_total_saved(sp, cat_id))


def get_cat_progress(sp, cat_id):
    cat = next((c for c in sp.get("categories", []) if c["id"] == cat_id), None)
    if not cat or cat["target"] <= 0:
        return 0.0
    return min(1.0, get_cat_total_saved(sp, cat_id) / cat["target"])


def get_cat_balance(sp, cat_id, monthly_data=None):
    saved = get_cat_total_saved(sp, cat_id)
    if monthly_data:
        spent = get_savings_spent(monthly_data, cat_id)
    else:
        spent = sp.get("spent", {}).get(cat_id, 0.0)
    return saved - spent


def get_cat_rollover(sp, cat_id, monthly_data=None):
    cat = next((c for c in sp.get("categories", []) if c["id"] == cat_id), None)
    if not cat or not cat.get("next_year_target"):
        return 0.0
    balance = get_cat_balance(sp, cat_id, monthly_data)
    if balance > 0 and get_cat_progress(sp, cat_id) >= 1.0:
        return balance
    return 0.0


def build_savings_spent_cache():
    """Single query to get total spent per savings category across all months."""
    conn = get_connection()
    rows = conn.execute(
        "SELECT savings_category_id, SUM(amount) as total FROM expenses WHERE savings_category_id IS NOT NULL GROUP BY savings_category_id"
    ).fetchall()
    return {r["savings_category_id"]: r["total"] for r in rows}


def get_cascade_adjustments(sp, cat_id, spent_cache=None):
    """Surplus cascade deduction for post-deadline months."""
    cat = next((c for c in sp.get("categories", []) if c["id"] == cat_id), None)
    if not cat or not cat.get("next_year_target"):
        return {}
    deadline = cat.get("deadline_month", 12)
    if deadline >= 12:
        return {}
    if get_cat_progress(sp, cat_id) < 1.0:
        return {}
    actual_row = sp.get("actual_grid", {}).get(cat_id, {})
    total_saved = sum(actual_row.get(f"{m:02d}", 0.0) for m in range(1, deadline + 1))
    total_spent = spent_cache.get(cat_id, 0.0) if spent_cache is not None else _get_all_months_savings_spent(cat_id)
    if total_spent <= 0:
        return {}
    surplus = total_saved - total_spent
    if surplus <= 0:
        return {}
    grid_row = sp.get("grid", {}).get(cat_id, {})
    adjustments = {}
    remaining_surplus = surplus
    for m in range(deadline + 1, 13):
        if remaining_surplus <= 0:
            break
        mk = f"{m:02d}"
        planned = grid_row.get(mk, 0.0)
        if planned <= 0:
            continue
        if remaining_surplus >= planned:
            adjustments[mk] = 0.0
            remaining_surplus -= planned
        else:
            adjustments[mk] = planned - remaining_surplus
            remaining_surplus = 0.0
    return adjustments


def _get_all_months_savings_spent(savings_cat_id):
    """Sum expenses funded from a savings category (SQL version)."""
    conn = get_connection()
    row = conn.execute(
        "SELECT COALESCE(SUM(amount), 0.0) FROM expenses WHERE savings_category_id = ?",
        (savings_cat_id,)
    ).fetchone()
    return row[0]


# --- Expense Categories CRUD ---
def get_expense_categories(data):
    return data.get("expense_categories", [])


def add_expense_category(data, name):
    cat = {"id": generate_id(), "name": name}
    data.setdefault("expense_categories", []).append(cat)
    return cat


def edit_expense_category(data, cat_id, new_name):
    for cat in data.get("expense_categories", []):
        if cat["id"] == cat_id:
            cat["name"] = new_name
            return cat
    return None


def delete_expense_category(data, cat_id):
    data["expense_categories"] = [c for c in data.get("expense_categories", []) if c["id"] != cat_id]
    for exp in data.get("expenses", []):
        if exp.get("expense_category_id") == cat_id:
            exp["expense_category_id"] = None


# --- Expense CRUD ---
def add_expense(data, date, amount, split_category_id, expense_category_id=None, description="", tags=None):
    exp = {
        "id": generate_id(), "date": date, "amount": amount,
        "category_id": split_category_id, "expense_category_id": expense_category_id,
        "description": description, "tags": tags or [],
    }
    data.setdefault("expenses", []).append(exp)
    return exp


def edit_expense(data, exp_id, **fields):
    for exp in data.get("expenses", []):
        if exp["id"] == exp_id:
            exp.update(fields)
            return exp
    return None


def delete_expense(data, exp_id):
    data["expenses"] = [e for e in data.get("expenses", []) if e["id"] != exp_id]


# --- Cashflow ---
def get_cashflow(data):
    return data.setdefault("cashflow", {"current_accounts": {}, "shared_actual": 0.0, "shared_assumed": 0.0, "saved_assumed": 0.0})


# --- Summary Helpers ---
def get_expenses_by_expense_category(data):
    grouped = {}
    for exp in data.get("expenses", []):
        key = exp.get("expense_category_id") or "uncategorized"
        grouped.setdefault(key, []).append(exp)
    return grouped


def get_expense_category_total(data, expense_cat_id):
    return sum(e["amount"] for e in data.get("expenses", []) if e.get("expense_category_id") == expense_cat_id)


def get_expense_category_split_breakdown(data, expense_cat_id):
    breakdown = {}
    for exp in data.get("expenses", []):
        if exp.get("expense_category_id") == expense_cat_id:
            if exp.get("savings_category_id"):
                key = f"sav:{exp['savings_category_id']}"
            else:
                key = exp.get("category_id") or "unknown"
            breakdown[key] = breakdown.get(key, 0.0) + exp["amount"]
    return breakdown


def get_net_income(data):
    total = 0.0
    for item in data.get("income_items", []):
        if item["type"] == "addition":
            total += item["amount"]
        else:
            total -= item["amount"]
    return total


def get_allocated_amount(data, category_id):
    net_income = get_net_income(data)
    cat_extra = sum(item["amount"] for item in data.get("income_items", []) if item.get("category_id") == category_id and item["type"] == "addition")
    general_net = net_income - sum(item["amount"] for item in data.get("income_items", []) if item.get("category_id") and item["type"] == "addition")
    for cat in data.get("categories", []):
        if cat["id"] == category_id:
            return general_net * (cat["percent"] / 100.0) + cat_extra
    return 0.0


def get_spent_amount(data, category_id):
    total = 0.0
    for exp in data.get("expenses", []):
        if exp.get("category_id") == category_id and not exp.get("savings_category_id"):
            total += exp["amount"]
    return total


def get_savings_spent(data, savings_cat_id):
    total = 0.0
    for exp in data.get("expenses", []):
        if exp.get("savings_category_id") == savings_cat_id:
            total += exp["amount"]
    return total


def get_total_expenses(data):
    return sum(exp["amount"] for exp in data.get("expenses", []))


def get_remaining_amount(data, category_id):
    return get_allocated_amount(data, category_id) - get_spent_amount(data, category_id)


# --- xlsx import (stub) ---
def load_xlsx(filepath: str, password: str = None, sheet_name: str = "kwiecień 2026") -> dict:
    try:
        with open(filepath, 'rb') as f:
            if password:
                office_file = msoffcrypto.OfficeFile(f)
                office_file.load_key(password=password)
                decrypted = io.BytesIO()
                office_file.decrypt(decrypted)
                wb = openpyxl.load_workbook(decrypted, data_only=True)
            else:
                wb = openpyxl.load_workbook(f, data_only=True)
        sheet = None
        for name in wb.sheetnames:
            if name.strip().lower() == sheet_name.strip().lower():
                sheet = wb[name]
                break
        if not sheet:
            return {"error": "sheet_not_found", "available_sheets": wb.sheetnames}
        return {"status": "success", "message": "Imported basic layout"}
    except Exception as e:
        return {"error": "import_failed", "message": str(e)}
